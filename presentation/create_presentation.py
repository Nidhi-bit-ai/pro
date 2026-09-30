from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT_FILE = Path(__file__).resolve().parent / "MNIT_RAG_Presentation.pptx"

THEME = {
    "primary": RGBColor(15, 60, 110),
    "secondary": RGBColor(31, 119, 180),
    "accent": RGBColor(0, 153, 153),
    "text": RGBColor(33, 37, 41),
    "muted": RGBColor(108, 117, 125),
    "bg": RGBColor(245, 248, 252),
}


def add_title_banner(slide, title: str, subtitle: str | None = None) -> None:
    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bg.line.fill.background()

    banner = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(1.0))
    banner.fill.solid()
    banner.fill.fore_color.rgb = THEME["primary"]
    banner.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.18), Inches(9.8), Inches(0.6))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(9.2), Inches(0.25), Inches(3.8), Inches(0.5))
        sp = sub_box.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.alignment = PP_ALIGN.RIGHT
        sp.font.size = Pt(12)
        sp.font.color.rgb = RGBColor(230, 240, 255)


def add_bullets(slide, left: float, top: float, width: float, height: float, bullets: list[str], level0_size: int = 20) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()

    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if bullet.startswith("  - "):
            p.text = bullet[4:]
            p.level = 1
            p.font.size = Pt(level0_size - 2)
        else:
            p.text = bullet
            p.level = 0
            p.font.size = Pt(level0_size)
            p.font.bold = True
        p.font.color.rgb = THEME["text"]


def add_footer(slide, text: str) -> None:
    footer = slide.shapes.add_textbox(Inches(0.5), Inches(7.1), Inches(12.3), Inches(0.3))
    p = footer.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(11)
    p.font.color.rgb = THEME["muted"]


def title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_banner(slide, "MNIT Document RAG System", "Repository: Nidhi-bit-ai/pro")

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.8), Inches(1.2))
    p = title_box.text_frame.paragraphs[0]
    p.text = "Implementation-grounded project presentation"
    p.font.size = Pt(34)
    p.font.bold = True
    p.font.color.rgb = THEME["primary"]

    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.3), Inches(12.0), Inches(2.0))
    sub_tf = sub_box.text_frame
    sub_tf.word_wrap = True
    lines = [
        "Covers scraping, metadata indexing, hybrid retrieval, and API deployment.",
        "Prepared from repository code and notebooks (RAG/, RAG_Server/).",
    ]
    for i, line in enumerate(lines):
        p = sub_tf.paragraphs[0] if i == 0 else sub_tf.add_paragraph()
        p.text = line
        p.font.size = Pt(20)
        p.font.color.rgb = THEME["text"]

    add_footer(slide, "MNIT Jaipur knowledge access via Retrieval-Augmented Generation (RAG)")


def slide_problem(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_banner(slide, "1) Introduction of the problem")

    bullets = [
        "MNIT information is spread across many PDFs and web pages.",
        "  - Students/faculty need fast answers for syllabus, fees, policies, notices, programs.",
        "Manual search is slow and error-prone when documents are large/unstructured.",
        "Need: an accurate question-answering pipeline grounded in official MNIT sources.",
    ]
    add_bullets(slide, 0.8, 1.4, 12.0, 4.6, bullets, level0_size=21)
    add_footer(slide, "Evidence: RAG notebooks and FastAPI routers in RAG_Server")


def slide_applications(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_banner(slide, "2) Applications")

    bullets = [
        "Academic support and policy lookup",
        "  - First-year Mathematics-I syllabus and curriculum extraction",
        "  - Fee structure discovery by academic year",
        "Administrative intelligence",
        "  - UFM policy, RTI procedure, PhD exam guidelines retrieval",
        "Research and faculty discovery",
        "  - Queries on faculty programs (e.g., quantum computing, smart-grid)",
        "Student analytics",
        "  - Toppers/achievement-related question answering from indexed notices",
    ]
    add_bullets(slide, 0.8, 1.35, 12.0, 5.4, bullets, level0_size=19)
    add_footer(slide, "Examples demonstrated in RAG/retrieval/retrieval.ipynb")


def slide_objectives(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_banner(slide, "3) Objectives")

    bullets = [
        "Ingest MNIT web content and automatically collect PDF knowledge sources.",
        "Generate structured metadata to improve filtering and traceability.",
        "Build robust retrieval by combining dense vectors + keyword BM25 search.",
        "Improve answer quality with multi-query expansion, RRF fusion, and LLM reranking.",
        "Expose production-friendly workflows through /scrapping, /indexing, /retrieval APIs.",
    ]
    add_bullets(slide, 0.8, 1.5, 12.0, 4.9, bullets, level0_size=20)
    add_footer(slide, "FastAPI entrypoint: RAG_Server/main.py")


def slide_io_eval(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_banner(slide, "4) Input, Output and Evaluation Parameters")

    table = slide.shapes.add_table(rows=5, cols=3, left=Inches(0.7), top=Inches(1.5), width=Inches(12.0), height=Inches(4.2)).table
    table.columns[0].width = Inches(2.6)
    table.columns[1].width = Inches(4.8)
    table.columns[2].width = Inches(4.6)

    headers = ["Component", "Implementation-backed details", "Evaluation (targets / measurement plan)"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(13)

    rows = [
        ("Input", "MNIT website pages and linked PDFs crawled from mnit.ac.in", "Track crawl coverage: pages visited, PDFs downloaded"),
        ("Indexed Corpus", "Chroma vector store: 2,339 docs; BM25 index: 2,339 docs (retrieval notebook run)", "Monitor index freshness and deduplication across re-index runs"),
        ("Output", "Top-N retrieved chunks + source metadata + reranker reasoning/score", "Validate citation correctness and answer grounding"),
        ("Retrieval Quality", "Hybrid search: embeddings + BM25 + RRF + LLM rerank", "Plan: Recall@k, MRR/NDCG, human relevance judgments on query set"),
    ]

    for r, row in enumerate(rows, start=1):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = text
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(12)

    add_footer(slide, "No benchmark scores claimed here; only observed counts and evaluation plan")


def slide_methodology(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_banner(slide, "5) Methodology: LLM algorithm + flow diagram")

    # Left: algorithm steps
    algo_box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.35), Inches(6.2), Inches(5.4))
    algo_box.fill.solid()
    algo_box.fill.fore_color.rgb = THEME["bg"]
    algo_box.line.color.rgb = THEME["secondary"]

    algo_text = algo_box.text_frame
    algo_text.word_wrap = True
    algo_lines = [
        "Algorithm (implementation mapping)",
        "1. Crawl mnit.ac.in and download PDFs (threaded crawler).",
        "2. Load PDFs (PyPDFLoader), extract metadata (Gemini + Pydantic parser).",
        "3. Chunk text (RecursiveCharacterTextSplitter).",
        "4. Embed chunks (Google text-embedding-004) and store in ChromaDB.",
        "5. For each user query: generate query variants with Gemini.",
        "6. Run vector similarity + BM25 keyword retrieval.",
        "7. Fuse ranked lists with Reciprocal Rank Fusion (RRF).",
        "8. Rerank candidates with LLM relevance scoring and return top sources.",
    ]
    for i, line in enumerate(algo_lines):
        p = algo_text.paragraphs[0] if i == 0 else algo_text.add_paragraph()
        p.text = line
        p.font.size = Pt(13 if i else 15)
        p.font.bold = i == 0
        p.font.color.rgb = THEME["text"]

    # Right: flow diagram
    x = 7.2
    y = 1.55
    w = 5.3
    h = 0.65
    steps = [
        "Web Crawl + PDF Download",
        "Metadata + Chunking + Embedding",
        "ChromaDB + BM25 Index",
        "User Query + Multi-Query",
        "Hybrid Retrieve + RRF + LLM Rerank",
        "Grounded Top-N Results",
    ]

    for i, label in enumerate(steps):
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y + i * 0.85), Inches(w), Inches(h))
        box.fill.solid()
        box.fill.fore_color.rgb = THEME["secondary"] if i % 2 == 0 else THEME["accent"]
        box.line.fill.background()
        tf = box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = label
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)

        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.DOWN_ARROW, Inches(x + 2.3), Inches(y + i * 0.85 + 0.64), Inches(0.7), Inches(0.20))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = THEME["primary"]
            arrow.line.fill.background()

    add_footer(slide, "Core modules: RAG_Server/src/scrapping, indexing, retrieval")


def slide_references(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_banner(slide, "6) Reference (Literature Survey)")

    refs = [
        "Lewis et al. (2020). Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks.",
        "Karpukhin et al. (2020). Dense Passage Retrieval for Open-Domain Question Answering.",
        "Robertson & Zaragoza (2009). The Probabilistic Relevance Framework: BM25 and Beyond.",
        "Cormack, Clarke & Buettcher (2009). Reciprocal Rank Fusion outperforms Condorcet and learning methods.",
        "LangChain documentation (retrievers/chains), Chroma vector store documentation.",
        "Google Generative AI embeddings/chat model references used in implementation.",
    ]

    add_bullets(slide, 0.8, 1.5, 12.0, 4.8, refs, level0_size=17)
    note = slide.shapes.add_textbox(Inches(0.8), Inches(6.2), Inches(12.0), Inches(0.7))
    p = note.text_frame.paragraphs[0]
    p.text = "Repository implementation references: RAG notebooks + RAG_Server/src services."
    p.font.size = Pt(13)
    p.font.color.rgb = THEME["muted"]
    add_footer(slide, "This slide mixes foundational papers with stack-specific implementation references")


def build_presentation(output_file: Path = OUTPUT_FILE) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.33)  # 16:9
    prs.slide_height = Inches(7.5)

    title_slide(prs)
    slide_problem(prs)
    slide_applications(prs)
    slide_objectives(prs)
    slide_io_eval(prs)
    slide_methodology(prs)
    slide_references(prs)

    prs.save(output_file)
    return output_file


if __name__ == "__main__":
    out = build_presentation()
    print(f"Presentation generated: {out}")
